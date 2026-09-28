from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF
from PIL import Image

from backend.utils.sanitize import sanitize_text, split_terms

ASSET_DIR = Path(__file__).resolve().parents[2] / "assets"
DEFAULT_LOGO = ASSET_DIR / "logo.png"


def _safe_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return value[:80] or "legal_document"


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def _add_docx_header(doc: Document, logo_path: str | Path | None, title: str) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if logo_path and Path(logo_path).exists():
        run = p.add_run()
        run.add_picture(str(logo_path), width=Inches(0.75))
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)


def format_docx(text: str, doc_type: str, terms: str = "", company_name: str = "LegalEase", footer_text: str = "Generated with LegalEase", logo_path: str | Path | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles["Normal"]
    styles.font.name = "Times New Roman"
    styles.font.size = Pt(11)

    logo = logo_path if logo_path else DEFAULT_LOGO
    _add_docx_header(doc, logo, doc_type)

    company = doc.add_paragraph()
    company.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = company.add_run(company_name)
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)

    # Preserve AI section structure while making common headings visually clear.
    for raw in sanitize_text(text).split("\n"):
        line = raw.strip()
        if not line:
            continue
        is_heading = bool(re.match(r"^(\d+\.?\s+|[A-Z][A-Z\s&/-]{4,}:?$)", line)) and len(line) < 120
        if is_heading:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line.rstrip(":"))
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15

    term_items = split_terms(terms)
    if term_items:
        doc.add_paragraph()
        h = doc.add_paragraph()
        r = h.add_run("KEY TERMS")
        r.bold = True
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "#"
        table.rows[0].cells[1].text = "Term / Condition"
        for idx, term in enumerate(term_items, 1):
            cells = table.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(f"{footer_text} | {company_name}")
    fr.font.name = "Times New Roman"
    fr.font.size = Pt(8)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, logo_path: str | Path | None, footer_text: str, company_name: str):
        super().__init__()
        self.logo_path = str(logo_path) if logo_path and Path(logo_path).exists() else None
        self.footer_text = footer_text
        self.company_name = company_name

    def header(self):
        if self.logo_path:
            try:
                self.image(self.logo_path, x=95, y=8, w=20)
            except Exception:
                pass
        self.set_x(10)
        self.set_y(30)

    def footer(self):
        self.set_y(-16)
        self.set_font("Helvetica", size=8)
        self.cell(0, 5, f"{self.footer_text} | {self.company_name} | Page {self.page_no()}", align="C")


def format_pdf(text: str, doc_type: str, company_name: str = "LegalEase", footer_text: str = "Generated with LegalEase", logo_path: str | Path | None = None) -> bytes:
    pdf = LegalEasePDF(logo_path or DEFAULT_LOGO, footer_text, company_name)
    pdf.set_auto_page_break(auto=True, margin=22)
    pdf.add_page()
    pdf.set_x(10)
    pdf.set_title(doc_type)
    pdf.set_author(company_name)

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, doc_type.upper(), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_x(10)
    pdf.multi_cell(0, 6, company_name, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    for raw in sanitize_text(text).split("\n"):
        line = raw.strip()
        if not line:
            pdf.ln(2)
            continue
        is_heading = bool(re.match(r"^(\d+\.?\s+|[A-Z][A-Z\s&/-]{4,}:?$)", line)) and len(line) < 120
        if is_heading:
            pdf.set_font("Helvetica", "B", 11)
            pdf.ln(2)
            pdf.multi_cell(0, 6, line.rstrip(":"), new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", size=10)
            pdf.multi_cell(0, 5.5, line, new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


def export_document(fmt: str, content: str, doc_type: str, terms: str = "", company_name: str = "LegalEase", footer_text: str = "Generated with LegalEase", logo_path: str | Path | None = None) -> tuple[bytes, str, str]:
    fmt = fmt.lower().lstrip(".")
    safe = _safe_filename(doc_type)
    if fmt == "txt":
        return format_txt(content), "text/plain; charset=utf-8", f"{safe}.txt"
    if fmt == "docx":
        return format_docx(content, doc_type, terms, company_name, footer_text, logo_path), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{safe}.docx"
    if fmt == "pdf":
        return format_pdf(content, doc_type, company_name, footer_text, logo_path), "application/pdf", f"{safe}.pdf"
    raise ValueError("Unsupported format. Use txt, docx, or pdf.")
